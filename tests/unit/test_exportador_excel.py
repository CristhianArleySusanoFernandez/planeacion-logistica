"""El exportador escribe las 4 hojas con el layout de la referencia y los totales cuadran."""

from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from planeacion.application.dto.pivote import ClienteNoResueltoDTO, FacturaDTO
from planeacion.application.dto.planeacion import PlaneacionCompleta
from planeacion.domain.modelo import (
    CargaCarro,
    Carro,
    MetricasDesbalance,
    MotivoNoResuelto,
    Municipio,
    ResultadoBalanceo,
    Zona,
    ZonaAgregada,
)
from planeacion.infraestructura.adaptadores.salida.excel.exportador_planeacion import (
    ENCABEZADO_BASE,
    ENCABEZADO_ECOM,
    ENCABEZADO_PEDIDOS,
    ENCABEZADO_PLANEACION,
    ExportadorExcelPlaneacion,
)

_BARBOSA = Municipio(nombre="BARBOSA")
_ZONA_UNO = Zona(nombre="(BARBOSA):  UNO", municipio=_BARBOSA)
_ZONA_DOS = Zona(nombre="(BARBOSA):  DOS", municipio=_BARBOSA)


def _factura(pedido: str, codigo: str, zona: str | None, total: str, kilos: str) -> FacturaDTO:
    return FacturaDTO(
        pedido=pedido,
        codigo_cliente=codigo,
        zona=zona,
        total=Decimal(total),
        kilos=Decimal(kilos),
        fecha=date(2026, 7, 8),
        nombre_cliente=f"TIENDA {codigo}",
        ciudad="BARBOSA",
        barrio="CENTRO",
        direccion="CL 1 2 3",
        asesor="10-ASESORA",
    )


def _planeacion_de_ejemplo() -> PlaneacionCompleta:
    carga_10 = CargaCarro(
        carro=Carro(numero="10", conductor="ANA", auxiliar="LUIS", municipio=_BARBOSA),
        zonas=[
            ZonaAgregada(
                zona=_ZONA_UNO, facturas=2, clientes=2, pesos=Decimal("1000.50"), kilos=Decimal("12.5")
            )
        ],
    )
    carga_20 = CargaCarro(
        carro=Carro(numero="20", conductor="PEDRO", municipio=_BARBOSA),
        zonas=[
            ZonaAgregada(zona=_ZONA_DOS, facturas=1, clientes=1, pesos=Decimal("500"), kilos=Decimal("3"))
        ],
    )
    metricas = MetricasDesbalance(
        municipio="BARBOSA",
        rango_clientes=1,
        rango_pesos=Decimal("500.50"),
        cv_clientes=0.1,
        cv_pesos=0.2,
    )
    resultado = ResultadoBalanceo(
        cargas_por_municipio={"BARBOSA": [carga_10, carga_20]},
        metricas_iniciales={"BARBOSA": metricas},
        metricas_finales={"BARBOSA": metricas},
        desde_historico=False,
    )
    facturas = (
        _factura("P1", "111", _ZONA_UNO.nombre, "600.50", "7.5"),
        _factura("P2", "222", _ZONA_UNO.nombre, "400", "5"),
        _factura("P3", "333", _ZONA_DOS.nombre, "500", "3"),
        _factura("P4", "999", None, "100", "1"),  # cliente sin zona resuelta
    )
    return PlaneacionCompleta(
        fecha=date(2026, 7, 8),
        dia_semana="miercoles",
        resultado=resultado,
        total_facturas=4,
        total_clientes=4,
        total_pesos=Decimal("1600.50"),
        total_kilos=Decimal("16.5"),
        no_resueltos=(
            ClienteNoResueltoDTO(
                codigo="999",
                motivo=MotivoNoResuelto.NO_ESTA_EN_MAESTRA.value,
                nombre="TIENDA 999",
                documento=None,
                direccion=None,
                ciudad=None,
                barrio=None,
            ),
        ),
        pedidos_excluidos_por_fecha=0,
        facturas=facturas,
        # Los agregados de lo no resuelto los calcula el dominio; el exportador
        # solo los escribe. Acá coinciden con la factura P4.
        facturas_no_resueltas=1,
        pesos_no_resueltos=Decimal("100"),
        kilos_no_resueltos=Decimal("1"),
    )


def _exportar(tmp_path: Path) -> Path:
    return ExportadorExcelPlaneacion().exportar(_planeacion_de_ejemplo(), tmp_path / "salida.xlsx")


def test_crea_las_cuatro_hojas(tmp_path: Path) -> None:
    libro = load_workbook(_exportar(tmp_path))
    assert libro.sheetnames == ["ECOM", "PLANEACION", "BASE", "PEDIDOS"]


def test_hoja_ecom_cliente_a_carro(tmp_path: Path) -> None:
    hoja = load_workbook(_exportar(tmp_path))["ECOM"]
    assert (hoja["A3"].value, hoja["B3"].value) == ENCABEZADO_ECOM
    filas = [(hoja.cell(row=f, column=1).value, hoja.cell(row=f, column=2).value) for f in range(4, 9)]
    # Orden: por carro y código; los sin zona al final como #N/A; cierra Total general.
    assert filas == [
        ("111", "10"),
        ("222", "10"),
        ("333", "20"),
        ("999", "#N/A"),
        ("Total general", 4),
    ]


def test_hoja_planeacion_replica_el_pivote(tmp_path: Path) -> None:
    hoja = load_workbook(_exportar(tmp_path))["PLANEACION"]
    assert hoja["B2"].value == "Promedio Vh"
    assert hoja["B3"].value == "Ventas Totales"
    assert (hoja["C3"].value, hoja["F3"].value) == (4, 4)
    assert hoja["D3"].value == 1600.50
    assert hoja["E3"].value == 16.5  # los totales en KILOS
    encabezados = tuple(hoja.cell(row=4, column=c).value for c in range(1, 7))
    assert encabezados == ENCABEZADO_PLANEACION

    # Zonas en orden alfabético; columna E en GRAMOS y G repite el carro.
    assert [hoja.cell(row=f, column=2).value for f in (5, 6, 7)] == [
        _ZONA_DOS.nombre,
        _ZONA_UNO.nombre,
        "#N/D",
    ]
    assert (hoja["A5"].value, hoja["C5"].value, hoja["E5"].value, hoja["G5"].value) == ("20", 1, 3000, "20")
    assert hoja["E6"].value == 12500
    # La fila #N/D acumula lo no resuelto.
    assert (hoja["C7"].value, hoja["D7"].value, hoja["E7"].value, hoja["F7"].value) == (1, 100, 1000, 1)


def test_fila_no_resueltos_usa_los_agregados_del_dominio(tmp_path: Path) -> None:
    """La fila '#N/D' se escribe con lo que trae el DTO, no recalculando el detalle.

    Los valores de abajo se apartan a propósito de lo que daría sumar las
    facturas sin zona: si el exportador volviera a agregarlas por su cuenta,
    escribiría (1, 100, 1000, 1) y la prueba fallaría. Así queda fijado que la
    única definición de "cuánto pesa lo no resuelto" vive en AgregadorPorZona.
    """
    planeacion = replace(
        _planeacion_de_ejemplo(),
        facturas_no_resueltas=7,
        pesos_no_resueltos=Decimal("777.25"),
        kilos_no_resueltos=Decimal("4"),
    )
    ruta = ExportadorExcelPlaneacion().exportar(planeacion, tmp_path / "salida.xlsx")
    hoja = load_workbook(ruta)["PLANEACION"]
    assert hoja["B7"].value == "#N/D"
    # E en gramos, como el resto de las filas de zona; F son los clientes distintos.
    assert (hoja["C7"].value, hoja["D7"].value, hoja["E7"].value, hoja["F7"].value) == (
        7,
        777.25,
        4000,
        1,
    )


def test_sin_clientes_no_resueltos_no_se_escribe_la_fila(tmp_path: Path) -> None:
    planeacion = replace(
        _planeacion_de_ejemplo(),
        no_resueltos=(),
        facturas_no_resueltas=0,
        pesos_no_resueltos=Decimal("0"),
        kilos_no_resueltos=Decimal("0"),
    )
    ruta = ExportadorExcelPlaneacion().exportar(planeacion, tmp_path / "salida.xlsx")
    hoja = load_workbook(ruta)["PLANEACION"]
    assert hoja["B7"].value is None


def test_hoja_base_resume_por_carro(tmp_path: Path) -> None:
    hoja = load_workbook(_exportar(tmp_path))["BASE"]
    assert tuple(hoja.cell(row=1, column=c).value for c in range(1, 9)) == ENCABEZADO_BASE
    fila_10 = tuple(hoja.cell(row=2, column=c).value for c in range(1, 10))
    assert fila_10 == ("10", 2, 2, 1000.50, 12.5, "ANA", "LUIS", "BARBOSA", _ZONA_UNO.nombre)
    assert hoja["A3"].value == "20"
    assert hoja["A4"].value is None  # solo carros con zonas


def test_hoja_pedidos_una_fila_por_factura(tmp_path: Path) -> None:
    hoja = load_workbook(_exportar(tmp_path))["PEDIDOS"]
    assert tuple(hoja.cell(row=1, column=c).value for c in range(1, 13)) == ENCABEZADO_PEDIDOS
    filas = [tuple(hoja.cell(row=f, column=c).value for c in range(1, 13)) for f in range(2, 6)]
    assert len(filas) == 4
    primera = filas[0]
    assert primera[1:8] == ("P1", "111", "TIENDA 111", "BARBOSA", "CENTRO", "CL 1 2 3", 600.50)
    assert primera[8:] == (7.5, "10-ASESORA", _ZONA_UNO.nombre, "10")
    # El cliente sin zona sale con #N/D y #N/A pero no se pierde.
    assert filas[3][10:] == ("#N/D", "#N/A")
    assert hoja.cell(row=6, column=1).value is None
