"""Flujo completo con el ECOM real: pivote + balanceo + exportación.

Verifica que el Excel de salida trae las 4 hojas, que el encabezado de la hoja
ECOM es idéntico al del .xlsm de referencia (leído del propio archivo), y que
los totales de la hoja PLANEACION reproducen los del pivote.
"""

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

RUTA_ECOM = Path("datos/pedidos24-26Junio.xlsx")
RUTA_REFERENCIA = Path("datos/DEL_24_PAR_EL_26_JUNIO.xlsm")
_TOLERANCIA = 0.01


def _hay_credenciales() -> bool:
    try:
        from planeacion.config.settings import Settings

        Settings()
        return True
    except ValidationError:
        return False


@pytest.mark.integration
@pytest.mark.skipif(
    not (_hay_credenciales() and RUTA_ECOM.exists() and RUTA_REFERENCIA.exists()),
    reason="sin credenciales de Supabase o sin los archivos de datos/",
)
def test_exportacion_del_flujo_completo(tmp_path: Path) -> None:
    from planeacion.config.contenedor import (
        crear_contenedor,
        crear_exportador_planeacion,
        crear_generar_planeacion,
    )

    planeacion = crear_generar_planeacion(crear_contenedor()).ejecutar(RUTA_ECOM)
    ruta_salida = crear_exportador_planeacion().exportar(planeacion, tmp_path / "salida.xlsx")

    libro = load_workbook(ruta_salida)
    assert libro.sheetnames == ["ECOM", "PLANEACION", "BASE", "PEDIDOS"]

    # El encabezado de la hoja ECOM debe ser el mismo del archivo de referencia.
    referencia = load_workbook(RUTA_REFERENCIA, read_only=True, data_only=True)
    try:
        hoja_referencia = referencia["ECOM"]
        encabezado_referencia = [
            tuple(fila) for fila in hoja_referencia.iter_rows(min_row=3, max_row=3, max_col=2,
                                                              values_only=True)
        ][0]
    finally:
        referencia.close()
    hoja_ecom = libro["ECOM"]
    assert (hoja_ecom["A3"].value, hoja_ecom["B3"].value) == encabezado_referencia

    # Una fila por cliente (los sin zona incluidos) y el cierre Total general.
    filas_clientes = 0
    fila = 4
    while hoja_ecom.cell(row=fila, column=1).value not in (None, "Total general"):
        filas_clientes += 1
        fila += 1
    assert filas_clientes == planeacion.total_clientes
    assert hoja_ecom.cell(row=fila, column=1).value == "Total general"

    # La hoja PLANEACION reproduce los totales del pivote.
    hoja_plan = libro["PLANEACION"]
    assert hoja_plan["B3"].value == "Ventas Totales"
    assert hoja_plan["C3"].value == planeacion.total_facturas
    assert abs(Decimal(str(hoja_plan["D3"].value)) - planeacion.total_pesos) <= Decimal("0.01")
    assert abs(Decimal(str(hoja_plan["E3"].value)) - planeacion.total_kilos) <= Decimal("0.01")
    assert hoja_plan["F3"].value == planeacion.total_clientes

    # Y la suma de facturas de sus filas (zonas + #N/D) también cuadra.
    suma_facturas = 0
    fila = 5
    while hoja_plan.cell(row=fila, column=3).value is not None:
        suma_facturas += int(hoja_plan.cell(row=fila, column=3).value)
        fila += 1
    assert suma_facturas == planeacion.total_facturas

    # PEDIDOS: una fila por factura.
    hoja_pedidos = libro["PEDIDOS"]
    filas_pedidos = 0
    fila = 2
    while hoja_pedidos.cell(row=fila, column=2).value is not None:
        filas_pedidos += 1
        fila += 1
    assert filas_pedidos == planeacion.total_facturas
