"""Los tres formatos que puede tener el ECOM dan exactamente los mismos totales.

Las pruebas unitarias del lector trabajan con un libro armado a mano; acá se usa
un archivo real de ECOM y se lo vuelve a guardar como .xlsm y como .xls, para que
lo que se compare sea el archivo del día y no un caso de laboratorio.
"""

from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from planeacion.domain.modelo import LineaPedido
from planeacion.infraestructura.adaptadores.salida.excel.lector_ecom import LectorEcomExcel

_ORIGINAL = Path("datos/pedidos24-26Junio.xlsx")


def _totales(lineas: list[LineaPedido]) -> tuple[int, int, Decimal, Decimal]:
    return (
        len(lineas),
        len({linea.pedido for linea in lineas}),
        sum((linea.total_linea for linea in lineas), Decimal("0")),
        sum((linea.kilos for linea in lineas), Decimal("0")),
    )


@pytest.fixture
def esperados() -> tuple[int, int, Decimal, Decimal]:
    if not _ORIGINAL.exists():
        pytest.skip(f"sin {_ORIGINAL}")
    return _totales(LectorEcomExcel().leer(_ORIGINAL))


@pytest.mark.integration
def test_el_mismo_archivo_como_xlsm_da_los_mismos_totales(
    tmp_path: Path, esperados: tuple[int, int, Decimal, Decimal]
) -> None:
    """Un .xlsm es el mismo OOXML con macros: los bytes valen tal cual."""
    copia = tmp_path / "ecom.xlsm"
    copia.write_bytes(_ORIGINAL.read_bytes())

    assert _totales(LectorEcomExcel().leer(copia)) == esperados


@pytest.mark.integration
def test_el_mismo_archivo_como_xls_binario_da_los_mismos_totales(
    tmp_path: Path, esperados: tuple[int, int, Decimal, Decimal]
) -> None:
    """Round-trip completo por el formato viejo: openpyxl lee el .xlsx original,
    xlwt lo escribe como BIFF8 y el lector lo recupera por el camino de xlrd."""
    xlwt = pytest.importorskip("xlwt")

    origen = openpyxl.load_workbook(_ORIGINAL, read_only=True, data_only=True)
    try:
        hoja_origen = origen[origen.sheetnames[0]]
        libro = xlwt.Workbook()
        hoja = libro.add_sheet("Hoja1")
        estilo_fecha = xlwt.easyxf(num_format_str="YYYY-MM-DD")
        for numero, fila in enumerate(hoja_origen.iter_rows(values_only=True)):
            for columna, valor in enumerate(fila):
                if valor is None:
                    continue
                if isinstance(valor, datetime):
                    hoja.write(numero, columna, valor, estilo_fecha)
                else:
                    hoja.write(numero, columna, valor)
    finally:
        origen.close()

    copia = tmp_path / "ecom.xls"
    libro.save(str(copia))

    assert _totales(LectorEcomExcel().leer(copia)) == esperados
