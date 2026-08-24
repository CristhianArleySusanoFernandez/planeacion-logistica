"""Lector del bloque de ECOM que viene pegado dentro de un .xlsm de planeación.

Cada .xlsm de un día trae, en la hoja ``PEDIDOS`` y a partir de la columna S, una
copia literal del export crudo de ECOM de ese día (mismo encabezado, mismas dos
columnas "Total", kilos en gramos). Eso permite reprocesar planeaciones viejas
sin conservar los .xlsx sueltos, que es lo que hace ``planeacion-validar``.

Por qué se rebana desde la S y no se mapea la fila 1 completa: a la izquierda del
bloque (columnas A-P) hay OTRA tabla, la de trabajo manual, con encabezados
homónimos —``Fecha``, ``Kilos`` y dos ``Total`` más—. Mapeando la fila entera
habría cuatro columnas "Total" y el ``Kilos`` de la izquierda ganaría el match por
estar primero: el archivo se leería mal en silencio. Anclando el bloque en su
propio ``Tipo``/``Pedido`` se reutiliza tal cual el parseo de ``lector_ecom``.
"""

from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from planeacion.domain.modelo import LineaPedido
from planeacion.infraestructura.adaptadores.salida.excel.lector_ecom import (
    ColumnasEcomFaltantes,
    leer_bloque_ecom,
    normalizar_encabezado,
)

HOJA_PEDIDOS = "PEDIDOS"

# Primera columna (0-based) donde puede empezar el bloque: S. Antes de ella está
# la tabla de trabajo manual, que nunca debe entrar al mapeo.
_PRIMERA_COLUMNA_POSIBLE = 18

# El bloque se ancla en la primera de estas columnas: son las dos con las que
# arranca el export de ECOM y ninguna existe en la tabla de la izquierda.
_ENCABEZADOS_ANCLA = ("tipo", "pedido")


def _localizar_bloque(hoja: Worksheet, nombre_archivo: str) -> int:
    """Índice 0-based de la columna donde arranca el bloque de ECOM."""
    fila_1 = next(hoja.iter_rows(min_row=1, max_row=1, values_only=True), ())
    for indice in range(_PRIMERA_COLUMNA_POSIBLE, len(fila_1)):
        if normalizar_encabezado(fila_1[indice]) in _ENCABEZADOS_ANCLA:
            return indice
    raise ColumnasEcomFaltantes(
        f"La hoja {HOJA_PEDIDOS} de {nombre_archivo} no trae el bloque de ECOM: no encontré "
        f"un encabezado 'Tipo' ni 'Pedido' de la columna S en adelante. "
        "¿Es un .xlsm de planeación con los pedidos crudos pegados?"
    )


class LectorEcomEmbebido:
    """Adaptador del puerto ``LectorDePedidos`` para los .xlsm de planeación."""

    def leer(self, ruta: Path) -> list[LineaPedido]:
        libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        try:
            if HOJA_PEDIDOS not in libro.sheetnames:
                raise ColumnasEcomFaltantes(
                    f"{ruta.name} no tiene hoja {HOJA_PEDIDOS} "
                    f"(trae: {', '.join(libro.sheetnames) or 'ninguna'})."
                )
            hoja = libro[HOJA_PEDIDOS]
            return leer_bloque_ecom(hoja, _localizar_bloque(hoja, ruta.name))
        finally:
            libro.close()
