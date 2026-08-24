"""Lector del .xlsm de referencia (hojas MAESTRA, CAMBIOS, "martha ojo" y BASE).

Los archivos pesan ~30 MB, por eso se abren con read_only=True. Las hojas
declaran dimensiones falsas (MAESTRA dice tener 1.048.576 filas), así que la
lectura se corta tras una racha de filas sin código.
"""

from collections.abc import Iterator
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from types import TracebackType
from typing import Any

import openpyxl

HOJA_MAESTRA = "MAESTRA"
HOJA_CAMBIOS = "CAMBIOS"
HOJA_OVERRIDES = "martha ojo"
HOJA_BASE = "BASE"
HOJA_PLANEACION = "PLANEACION"

HOJAS_DE_SIEMBRA = (HOJA_MAESTRA, HOJA_CAMBIOS, HOJA_OVERRIDES, HOJA_BASE)

# Columna (índice 0-based) → encabezado esperado en la fila 1. "martha ojo" no
# tiene encabezado, por eso no aparece aquí. El de BASE es el del bloque 1
# (vehículos físicos), que no se siembra pero sirve para confirmar que el
# archivo es el .xlsm de planeación esperado; la flota sale del bloque 2
# (ver leer_rutas).
ENCABEZADOS_ESPERADOS: dict[str, dict[int, str]] = {
    HOJA_MAESTRA: {1: "Direccion", 2: "Codigo_Postal", 3: "Ciudad", 5: "RUTA"},
    HOJA_CAMBIOS: {0: "Codigo del cliente", 5: "Ciudad Real", 6: "Barrio Real"},
    HOJA_BASE: {0: "Codigo", 1: "CONDUCTOR", 2: "Placa", 3: "AUX ENTREGA", 4: "ZONA"},
}

# En "martha ojo" el código va en A y la zona forzada en H (I es un duplicado).
_COL_OVERRIDE_CODIGO = 0
_COL_OVERRIDE_ZONA = 7

_MAX_FILAS_VACIAS_SEGUIDAS = 100


@dataclass(frozen=True)
class FilaMaestra:
    codigo: str
    direccion: str | None
    codigo_postal: str | None
    ciudad: str | None
    barrio: str | None
    ruta: str | None


@dataclass(frozen=True)
class FilaCambio:
    codigo: str
    ciudad_real: str | None
    barrio_real: str | None


@dataclass(frozen=True)
class FilaOverride:
    codigo: str
    zona: str


@dataclass(frozen=True)
class FilaPlaneacionZona:
    """Un par carro→zona de la hoja PLANEACION (col A = carro, col B = zona):
    la materia prima del repertorio de zonas por carro."""

    carro: str
    zona: str


@dataclass(frozen=True)
class TotalesPlaneacion:
    """La fila "Ventas Totales" de la hoja PLANEACION: el control de que el
    origen se leyó bien (deben cuadrar con los totales que calcula la app)."""

    facturas: int
    pesos: Decimal
    kilos: Decimal
    clientes: int


@dataclass(frozen=True)
class FilaRutaBase:
    """Bloque 2 de BASE (encabezado 'Ruta | Facturas | ...'): las rutas de
    reparto 1..N que usan Rudy y facturación. Esta es la flota que se siembra."""

    numero: str
    facturas: int
    conductor: str | None
    auxiliar: str | None
    ciudad: str | None


def _texto(valor: Any) -> str | None:
    """Celda → texto limpio. Los enteros que Excel guarda como float no llevan '.0'."""
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    return texto or None


class LectorReferenciaExcel:
    def __init__(self, ruta: Path) -> None:
        self._ruta = ruta
        self._libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)

    def __enter__(self) -> "LectorReferenciaExcel":
        return self

    def __exit__(
        self,
        tipo: type[BaseException] | None,
        valor: BaseException | None,
        traza: TracebackType | None,
    ) -> None:
        self.cerrar()

    def cerrar(self) -> None:
        self._libro.close()

    def hojas_faltantes(self) -> list[str]:
        return [hoja for hoja in HOJAS_DE_SIEMBRA if hoja not in self._libro.sheetnames]

    def primeras_filas(self, hoja: str, cuantas: int = 3, max_col: int = 9) -> list[list[str]]:
        """Primeras filas de una hoja, como texto, para imprimir la inspección."""
        filas: list[list[str]] = []
        for fila in self._libro[hoja].iter_rows(max_row=cuantas, max_col=max_col, values_only=True):
            filas.append([_texto(celda) or "" for celda in fila])
        return filas

    def errores_de_encabezado(self) -> list[str]:
        """Compara la fila 1 de cada hoja contra ENCABEZADOS_ESPERADOS."""
        errores: list[str] = []
        for hoja, esperados in ENCABEZADOS_ESPERADOS.items():
            fila_1 = self.primeras_filas(hoja, cuantas=1)[0]
            for indice, esperado in esperados.items():
                real = fila_1[indice] if indice < len(fila_1) else ""
                if real.strip().upper() != esperado.strip().upper():
                    errores.append(f"{hoja}: columna {indice + 1} esperaba '{esperado}' y trae '{real}'")
        return errores

    def _filas_con_codigo(
        self, hoja: str, max_col: int, min_row: int = 2
    ) -> Iterator[tuple[str, tuple[Any, ...]]]:
        """Itera (codigo, fila) saltando el encabezado, cortando tras una racha de filas sin código."""
        vacias_seguidas = 0
        for fila in self._libro[hoja].iter_rows(min_row=min_row, max_col=max_col, values_only=True):
            codigo = _texto(fila[0]) if fila else None
            if codigo is None:
                vacias_seguidas += 1
                if vacias_seguidas > _MAX_FILAS_VACIAS_SEGUIDAS:
                    return
                continue
            vacias_seguidas = 0
            yield codigo, fila

    def leer_maestra(self) -> Iterator[FilaMaestra]:
        for codigo, fila in self._filas_con_codigo(HOJA_MAESTRA, max_col=6):
            yield FilaMaestra(
                codigo=codigo,
                direccion=_texto(fila[1]),
                codigo_postal=_texto(fila[2]),
                ciudad=_texto(fila[3]),
                barrio=_texto(fila[4]),
                ruta=_texto(fila[5]),
            )

    def leer_cambios(self) -> Iterator[FilaCambio]:
        for codigo, fila in self._filas_con_codigo(HOJA_CAMBIOS, max_col=7):
            yield FilaCambio(
                codigo=codigo,
                ciudad_real=_texto(fila[5]),
                barrio_real=_texto(fila[6]),
            )

    def leer_overrides(self) -> Iterator[FilaOverride]:
        # "martha ojo" no tiene encabezado: la fila 1 ya es un dato.
        for codigo, fila in self._filas_con_codigo(HOJA_OVERRIDES, max_col=9, min_row=1):
            zona = _texto(fila[_COL_OVERRIDE_ZONA])
            if zona is None:
                continue
            yield FilaOverride(codigo=codigo, zona=zona)

    def tiene_hoja(self, hoja: str) -> bool:
        return hoja in self._libro.sheetnames

    def leer_asignaciones_planeacion(self) -> Iterator[FilaPlaneacionZona]:
        """Pares carro→zona de la hoja PLANEACION: col A = carro, col B = zona,
        datos desde la fila 5 (las filas 1-4 son totales y encabezados). Ignora
        las filas sin carro numérico (#N/D, totales) y las zonas sintéticas del
        pivote ('#N/D', '(en blanco)')."""
        no_asignables = {"#N/D", "#N/A", "(EN BLANCO)"}
        vacias_seguidas = 0
        hoja = self._libro[HOJA_PLANEACION]
        for fila in hoja.iter_rows(min_row=5, max_col=2, values_only=True):
            carro = _texto(fila[0]) if fila else None
            zona = _texto(fila[1]) if len(fila) > 1 else None
            if zona is None:
                vacias_seguidas += 1
                if vacias_seguidas > _MAX_FILAS_VACIAS_SEGUIDAS:
                    return
                continue
            vacias_seguidas = 0
            if carro is None or not carro.isdigit() or zona.strip().upper() in no_asignables:
                continue
            yield FilaPlaneacionZona(carro=carro, zona=zona)

    def leer_totales_planeacion(self) -> TotalesPlaneacion:
        """Fila 3 de PLANEACION ("Ventas Totales"): C facturas, D pesos, E kilos, F clientes.

        Ojo con los kilos: esta fila los trae en KILOS, mientras que las filas de
        zona de más abajo los traen en gramos (es el pivote viejo, que multiplica
        por mil al detallar). Los números salen de fórmulas de Excel, así que
        llegan como float y se pasan a Decimal por su texto para no arrastrar
        basura binaria; aun así conviene compararlos con tolerancia.
        """
        fila = next(self._libro[HOJA_PLANEACION].iter_rows(min_row=3, max_row=3, max_col=6, values_only=True))

        def numero(indice: int) -> Decimal:
            texto = _texto(fila[indice])
            return Decimal(texto) if texto else Decimal("0")

        return TotalesPlaneacion(
            facturas=int(numero(2)),
            pesos=numero(3),
            kilos=numero(4),
            clientes=int(numero(5)),
        )

    def leer_rutas(self) -> Iterator[FilaRutaBase]:
        """Lee el segundo bloque de BASE: arranca tras la fila-encabezado 'Ruta'
        (columnas A=Ruta B=Facturas F=CONDUCTOR G=AUX H=CIUDAD) y corta en la
        primera fila cuya primera celda no sea un número de ruta."""
        en_bloque = False
        for fila in self._libro[HOJA_BASE].iter_rows(max_col=8, values_only=True):
            primera = _texto(fila[0]) if fila else None
            if not en_bloque:
                en_bloque = primera is not None and primera.strip().upper() == "RUTA"
                continue
            if primera is None or not primera.isdigit():
                return
            facturas = _texto(fila[1])
            yield FilaRutaBase(
                numero=primera,
                facturas=int(facturas) if facturas and facturas.isdigit() else 0,
                conductor=_texto(fila[5]),
                auxiliar=_texto(fila[6]),
                ciudad=_texto(fila[7]),
            )
